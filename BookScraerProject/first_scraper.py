import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from urllib.parse import urljoin

url = "https://books.toscrape.com/"

wb = Workbook()
ws = wb.active
ws.title = "Books"

ws.append(["Title", "Price", "Rating", "Availability", "Product URL"])

session = requests.Session()
page_count = 0
book_count = 0

rating_map = {
    "One": 1,
    "Two": 2,
    "Three": 3,
    "Four": 4,
    "Five": 5
}

while url:
    print("正在抓取：", url)

    try:
        response = session.get(url, timeout=15)
        response.raise_for_status()
    except requests.RequestException as e:
        print("网页请求失败：", e)
        print("请检查网络连接，稍后重试。")
        break

    soup = BeautifulSoup(response.text, "html.parser")
    books = soup.select("article.product_pod")

    for book in books:
        title = book.h3.a.get("title")

        price_text = book.select_one(
            ".price_color"
        ).get_text(strip=True)
        price_text = price_text.replace("Â", "").replace("£", "")
        price = float(price_text)

        rating_classes = book.select_one(
            ".star-rating"
        ).get("class", [])

        rating = next(
            (rating_map[c] for c in rating_classes if c in rating_map),
            None
        )

        availability = book.select_one(
            ".availability"
        ).get_text(" ", strip=True)

        product_url = urljoin(url, book.h3.a.get("href"))

        ws.append([
            title, price, rating, availability, product_url
        ])

        row = ws.max_row
        ws.cell(row=row, column=2).number_format = '£0.00'

        link_cell = ws.cell(row=row, column=5)
        link_cell.hyperlink = product_url
        link_cell.style = "Hyperlink"

        book_count += 1

    page_count += 1

    next_button = soup.select_one("li.next a")

    if next_button:
        url = urljoin(url, next_button.get("href"))
    else:
        url = None

for cell in ws[1]:
    cell.font = Font(bold=True)

ws.freeze_panes = "A2"

for column_cells in ws.columns:
    max_length = max(
        len(str(cell.value or "")) for cell in column_cells
    )
    column_letter = get_column_letter(column_cells[0].column)
    ws.column_dimensions[column_letter].width = min(
        max_length + 2, 60
    )

wb.save("books.xlsx")

print("抓取完成！")
print("抓取页数：", page_count)
print("图书总数：", book_count)
print("文件已保存：books.xlsx")